"""
Projects view — Embu County Project Management System.
Professional filter bar, project details, and multi-format import
(Excel, Word, CSV, pasted text).
"""

import io
import re
from datetime import date, datetime

import pandas as pd
import streamlit as st

from utils.db import run_query, execute_write


# =====================================================
# HELPERS
# =====================================================
def _to_date(val):
    """Safely coerce a DB value to a Python date."""
    if val is None:
        return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, pd.Timestamp):
        return val.date()
    try:
        return datetime.fromisoformat(str(val)[:10]).date()
    except Exception:
        return None


def _fmt_kes(value):
    if value is None or pd.isna(value):
        return "KES 0"
    return f"KES {float(value):,.0f}"


def _status_badge(status):
    colors = {
        "Completed":    ("#065f46", "#d1fae5"),
        "In Progress":  ("#1e40af", "#dbeafe"),
        "Not Started":  ("#374151", "#e5e7eb"),
        "Stalled":      ("#92400e", "#fef3c7"),
        "Cancelled":    ("#991b1b", "#fee2e2"),
    }
    fg, bg = colors.get(status, ("#374151", "#e5e7eb"))
    return (
        f'<span style="background:{bg}; color:{fg}; padding:4px 12px; '
        f'border-radius:12px; font-size:12px; font-weight:600;">'
        f'{status}</span>'
    )


def _parse_date(value):
    """Parse a date from various formats."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    s = str(value).strip()
    if not s or s.lower() in ("nan", "none", ""):
        return None

    formats = [
        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d",
        "%d-%b-%y", "%d-%b-%Y", "%d %b %Y", "%d %B %Y",
        "%d.%m.%Y", "%Y.%m.%d",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue

    try:
        return pd.to_datetime(s, dayfirst=True).date()
    except Exception:
        return None


def _parse_number(value):
    """Parse '1,500,000', 'KES 1500000', etc. Returns None on failure."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    s = str(value).strip().replace(",", "")
    s = re.sub(r"(?i)kes", "", s).strip()
    if not s or s.lower() in ("nan", "none", ""):
        return 0.0

    try:
        return float(s)
    except ValueError:
        return None


def _clean_str(value):
    """Return a stripped string or None for empty/NaN values."""
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    s = str(value).strip()
    if not s or s.lower() in ("nan", "none"):
        return None
    return s


# =====================================================
# MAIN RENDER
# =====================================================
def render():
    st.title("📋 Project Management")
    st.caption("Embu County Government — Land of Opportunities")
    st.markdown("---")

    # ---- Load reference data ----
    try:
        subcounties = run_query("SELECT id, name FROM subcounties ORDER BY name")
        wards = run_query(
            "SELECT id, name, subcounty_id FROM wards ORDER BY name"
        )
        departments = run_query(
            "SELECT id, name FROM departments ORDER BY name"
        )
    except Exception as e:
        st.error(f"⚠️ Could not load reference data: {e}")
        return

    tab_view, tab_add, tab_import = st.tabs([
        "📊 View Projects",
        "➕ Add New Project",
        "📥 Import Projects",
    ])

    with tab_view:
        _render_projects_view(subcounties, wards, departments)

    with tab_add:
        _render_add_project(subcounties, wards, departments)

    with tab_import:
        _render_import_projects(subcounties, wards, departments)


# =====================================================
# TAB 1: VIEW PROJECTS
# =====================================================
def _render_projects_view(subcounties, wards, departments):
    try:
        projects_df = run_query("""
            SELECT
                p.id,
                p.project_name,
                s.name AS subcounty,
                w.name AS ward,
                d.name AS department,
                p.commencement_date,
                p.expected_completion_date,
                p.total_funding,
                p.total_expenditure,
                p.percentage_completion,
                p.project_status,
                p.remarks
            FROM projects p
            LEFT JOIN subcounties s ON p.subcounty_id = s.id
            LEFT JOIN wards w ON p.ward_id = w.id
            LEFT JOIN departments d ON p.department_id = d.id
            ORDER BY p.id DESC
        """)
    except Exception as e:
        st.error(f"⚠️ Could not load projects: {e}")
        return

    if projects_df.empty:
        st.info("No projects yet. Use the **Add New Project** or **Import Projects** tab.")
        return

    # =====================================================
    # PROFESSIONAL FILTER BAR
    # =====================================================
    st.markdown("### 🔎 Filter Projects")

    with st.container(border=True):
        row1_col1, row1_col2, row1_col3 = st.columns(3)

        sc_options = ["All Subcounties"] + sorted(
            projects_df['subcounty'].dropna().unique().tolist()
        )
        selected_subcounty = row1_col1.selectbox(
            "📍 Subcounty", sc_options, index=0,
            key="proj_filter_subcounty",
        )

        if selected_subcounty == "All Subcounties":
            ward_source = projects_df['ward'].dropna().unique().tolist()
        else:
            ward_source = projects_df[
                projects_df['subcounty'] == selected_subcounty
            ]['ward'].dropna().unique().tolist()
        ward_options = ["All Wards"] + sorted(ward_source)
        selected_ward = row1_col2.selectbox(
            "🗺️ Ward", ward_options, index=0,
            key="proj_filter_ward",
        )

        dept_options = ["All Departments"] + sorted(
            projects_df['department'].dropna().unique().tolist()
        )
        selected_dept = row1_col3.selectbox(
            "🏢 Department", dept_options, index=0,
            key="proj_filter_dept",
        )

        row2_col1, row2_col2, row2_col3 = st.columns(3)

        status_options = ["All Statuses"] + [
            "Completed", "In Progress", "Not Started", "Stalled", "Cancelled"
        ]
        selected_status = row2_col1.selectbox(
            "📊 Project Status", status_options, index=0,
            key="proj_filter_status",
        )

        search_term = row2_col2.text_input(
            "🔍 Search Project Name",
            placeholder="Type to search...",
            key="proj_filter_search",
        )

        row2_col3.markdown("<br>", unsafe_allow_html=True)
        if row2_col3.button(
            "🔄 Reset Filters",
            use_container_width=True,
            key="proj_filter_reset",
        ):
            for k in [
                "proj_filter_subcounty", "proj_filter_ward",
                "proj_filter_dept", "proj_filter_status",
                "proj_filter_search",
            ]:
                st.session_state.pop(k, None)
            st.rerun()

    # ---- Apply filters ----
    view_df = projects_df.copy()
    if selected_subcounty != "All Subcounties":
        view_df = view_df[view_df['subcounty'] == selected_subcounty]
    if selected_ward != "All Wards":
        view_df = view_df[view_df['ward'] == selected_ward]
    if selected_dept != "All Departments":
        view_df = view_df[view_df['department'] == selected_dept]
    if selected_status != "All Statuses":
        view_df = view_df[view_df['project_status'] == selected_status]
    if search_term.strip():
        view_df = view_df[
            view_df['project_name'].str.contains(
                search_term.strip(), case=False, na=False
            )
        ]

    # ---- Summary metrics ----
    st.markdown("### 📈 Summary")
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Projects", f"{len(view_df):,}")
    m2.metric(
        "Total Funding",
        _fmt_kes(view_df['total_funding'].sum() if not view_df.empty else 0),
    )
    m3.metric(
        "Total Expenditure",
        _fmt_kes(view_df['total_expenditure'].sum() if not view_df.empty else 0),
    )
    avg_pct = view_df['percentage_completion'].mean() if not view_df.empty else 0
    m4.metric("Avg. Completion", f"{avg_pct:.1f}%")

    st.markdown("---")

    # ---- Projects table ----
    st.markdown(f"### 📑 Projects ({len(view_df)})")

    if view_df.empty:
        st.warning("No projects match the current filters.")
        return

    display_df = view_df.copy()
    display_df.insert(0, 'S/No', range(1, len(display_df) + 1))
    display_df = display_df.rename(columns={
        'project_name': 'Project Name',
        'ward': 'Ward',
        'subcounty': 'Subcounty',
        'commencement_date': 'Commencement',
        'expected_completion_date': 'Expected Completion',
        'department': 'Department',
        'total_funding': 'Total Funding (KES)',
        'total_expenditure': 'Total Expenditure (KES)',
        'percentage_completion': '% Completion',
        'remarks': 'Remarks',
    })

    table_cols = [
        'S/No', 'Project Name', 'Ward', 'Subcounty',
        'Commencement', 'Expected Completion', 'Department',
        'Total Funding (KES)', 'Total Expenditure (KES)',
        '% Completion', 'Remarks',
    ]
    display_df = display_df[table_cols]

    st.dataframe(
        display_df, use_container_width=True, hide_index=True, height=420,
        column_config={
            "S/No": st.column_config.NumberColumn("S/No", width="small"),
            "Project Name": st.column_config.TextColumn(
                "Project Name", width="large"
            ),
            "Ward": st.column_config.TextColumn("Ward", width="small"),
            "Subcounty": st.column_config.TextColumn("Subcounty", width="small"),
            "Commencement": st.column_config.DateColumn(
                "Commencement", format="DD MMM YYYY", width="small"
            ),
            "Expected Completion": st.column_config.DateColumn(
                "Expected Completion", format="DD MMM YYYY", width="small"
            ),
            "Department": st.column_config.TextColumn("Department", width="medium"),
            "Total Funding (KES)": st.column_config.NumberColumn(
                "Total Funding (KES)", format="KES %,.0f"
            ),
            "Total Expenditure (KES)": st.column_config.NumberColumn(
                "Total Expenditure (KES)", format="KES %,.0f"
            ),
            "% Completion": st.column_config.ProgressColumn(
                "% Completion", format="%.1f%%",
                min_value=0, max_value=100,
            ),
            "Remarks": st.column_config.TextColumn("Remarks", width="medium"),
        },
    )

    # ---- Project details viewer ----
    st.markdown("---")
    st.markdown("### 🔍 Project Details")

    detail_options = view_df.copy()
    detail_options['label'] = (
        "#" + detail_options['id'].astype(str) + " — "
        + detail_options['project_name']
    )

    selected_label = st.selectbox(
        "Select a project to view full details",
        options=detail_options['label'].tolist(),
        key="proj_detail_select",
        label_visibility="collapsed",
    )

    selected_id = int(selected_label.split(" — ")[0].replace("#", ""))
    proj = view_df[view_df['id'] == selected_id].iloc[0]

    with st.container(border=True):
        h1, h2 = st.columns([3, 1])
        h1.markdown(f"#### {proj['project_name']}")
        h2.markdown(
            f"<div style='text-align:right;'>"
            f"{_status_badge(proj['project_status'])}</div>",
            unsafe_allow_html=True,
        )

        st.markdown("---")

        r1c1, r1c2, r1c3 = st.columns(3)
        r1c1.markdown(f"**📍 Subcounty**  \n{proj['subcounty'] or '—'}")
        r1c2.markdown(f"**🗺️ Ward**  \n{proj['ward'] or '—'}")
        r1c3.markdown(f"**🏢 Department**  \n{proj['department'] or '—'}")

        st.markdown("")

        r2c1, r2c2, r2c3 = st.columns(3)
        comm = _to_date(proj['commencement_date'])
        exp = _to_date(proj['expected_completion_date'])
        r2c1.markdown(
            f"**📅 Commencement Date**  \n"
            f"{comm.strftime('%d %b %Y') if comm else '—'}"
        )
        r2c2.markdown(
            f"**🎯 Expected Completion**  \n"
            f"{exp.strftime('%d %b %Y') if exp else '—'}"
        )
        r2c3.markdown(
            f"**⏱️ Duration**  \n"
            f"{(exp - comm).days if comm and exp else '—'}"
            f"{' days' if comm and exp else ''}"
        )

        st.markdown("")

        r3c1, r3c2, r3c3 = st.columns(3)
        funding = proj['total_funding'] or 0
        expenditure = proj['total_expenditure'] or 0
        balance = funding - expenditure
        r3c1.markdown(f"**💰 Total Funding**  \n{_fmt_kes(funding)}")
        r3c2.markdown(f"**💸 Total Expenditure**  \n{_fmt_kes(expenditure)}")
        r3c3.markdown(f"**🏦 Balance Remaining**  \n{_fmt_kes(balance)}")

        st.markdown("")

        pct = float(proj['percentage_completion'] or 0)
        st.markdown(f"**📊 Completion Progress — {pct:.1f}%**")
        st.progress(min(max(pct / 100.0, 0.0), 1.0))

        st.markdown("")
        st.markdown(
            f"**📝 Remarks**  \n{proj['remarks'] or '_No remarks recorded._'}"
        )

    # ---- CSV export ----
    st.markdown("---")
    csv = display_df.to_csv(index=False).encode('utf-8')
    st.download_button(
        "⬇️ Download Filtered Projects (CSV)",
        data=csv,
        file_name=f"embu_projects_{date.today().isoformat()}.csv",
        mime="text/csv",
    )

    st.markdown("---")

    with st.expander("✏️ Update Project"):
        _render_edit_form(projects_df, subcounties, wards, departments)

    with st.expander("🗑️ Delete Project"):
        _render_delete_form(projects_df)


# =====================================================
# TAB 2: ADD NEW PROJECT
# =====================================================
def _render_add_project(subcounties, wards, departments):
    if subcounties.empty or departments.empty:
        st.warning(
            "⚠️ Please add Subcounties and Departments in the **Settings** "
            "module before creating projects."
        )
        return

    with st.form("new_project", clear_on_submit=True):
        st.markdown("**Basic Details**")
        col1, col2 = st.columns(2)

        project_name = col1.text_input("Project Name *")

        subcounty_id = col2.selectbox(
            "Subcounty *",
            options=subcounties['id'].tolist(),
            format_func=lambda x: subcounties.loc[
                subcounties['id'] == x, 'name'
            ].values[0],
        )

        filtered_wards = wards[wards['subcounty_id'] == subcounty_id]
        if not filtered_wards.empty:
            ward_id = col1.selectbox(
                "Ward *",
                options=filtered_wards['id'].tolist(),
                format_func=lambda x: filtered_wards.loc[
                    filtered_wards['id'] == x, 'name'
                ].values[0],
            )
        else:
            col1.warning("No wards found for this subcounty.")
            ward_id = None

        department_id = col2.selectbox(
            "Department *",
            options=departments['id'].tolist(),
            format_func=lambda x: departments.loc[
                departments['id'] == x, 'name'
            ].values[0],
        )

        st.markdown("**Timeline**")
        col3, col4 = st.columns(2)
        commencement_date = col3.date_input("Project Commencement Date", value=None)
        expected_completion_date = col4.date_input(
            "Expected Date of Completion", value=None
        )

        st.markdown("**Financials (KES)**")
        col5, col6 = st.columns(2)
        total_funding = col5.number_input(
            "Total Funding", min_value=0.0, step=1000.0, format="%.2f"
        )
        total_expenditure = col6.number_input(
            "Total Expenditure", min_value=0.0, step=1000.0, format="%.2f"
        )

        st.markdown("**Progress**")
        col7, col8 = st.columns(2)
        percentage_completion = col7.slider(
            "Percentage (%) of Completion",
            min_value=0, max_value=100, value=0, step=1,
        )
        project_status = col8.selectbox(
            "Status",
            ["Not Started", "In Progress", "Completed", "Stalled", "Cancelled"],
        )

        remarks = st.text_area("Remarks")

        submitted = st.form_submit_button(
            "💾 Save Project", use_container_width=True
        )

        if submitted:
            if not project_name.strip() or not ward_id:
                st.error("Project Name and Ward are required.")
            elif total_expenditure > total_funding and total_funding > 0:
                st.error("Total Expenditure cannot exceed Total Funding.")
            elif (
                commencement_date and expected_completion_date
                and expected_completion_date < commencement_date
            ):
                st.error(
                    "Expected Completion Date cannot be before Commencement Date."
                )
            else:
                ok, msg = execute_write(
                    """
                    INSERT INTO projects
                        (project_name, subcounty_id, ward_id, department_id,
                         project_status, remarks, commencement_date,
                         expected_completion_date, total_funding,
                         total_expenditure, percentage_completion)
                    VALUES
                        (:name, :sc, :ward, :dept, :status, :remarks,
                         :comm, :exp, :fund, :spent, :pct)
                    """,
                    {
                        "name": project_name.strip(),
                        "sc": subcounty_id,
                        "ward": ward_id,
                        "dept": department_id,
                        "status": project_status,
                        "remarks": remarks.strip() or None,
                        "comm": commencement_date,
                        "exp": expected_completion_date,
                        "fund": total_funding,
                        "spent": total_expenditure,
                        "pct": float(percentage_completion),
                    },
                )
                if ok:
                    st.success("✅ Project saved successfully.")
                    st.cache_data.clear()
                    st.rerun()
                else:
                    st.error(f"❌ {msg}")


# =====================================================
# TAB 3: IMPORT PROJECTS (Excel, Word, CSV, paste)
# =====================================================
EXPECTED_COLUMNS = [
    "Project Name", "Subcounty", "Ward", "Department",
    "Commencement Date", "Expected Completion Date",
    "Total Funding", "Total Expenditure",
    "Percentage Completion", "Status", "Remarks",
]


def _render_import_projects(subcounties, wards, departments):
    st.markdown("### 📥 Bulk Import Projects")
    st.caption(
        "Upload from Excel, CSV, or Word — or paste data directly. "
        "Download a template below to see the required format."
    )

    # =====================================================
    # TEMPLATE DOWNLOADS
    # =====================================================
    with st.expander("📄 Download Templates", expanded=False):
        tcol1, tcol2, tcol3 = st.columns(3)
        template_df = pd.DataFrame(columns=EXPECTED_COLUMNS)

        # CSV
        tcol1.download_button(
            "⬇️ CSV Template",
            data=template_df.to_csv(index=False).encode("utf-8"),
            file_name="embu_projects_template.csv",
            mime="text/csv",
            use_container_width=True,
        )

        # Excel with reference sheet
        try:
            import openpyxl  # noqa: F401
            buf = io.BytesIO()
            with pd.ExcelWriter(buf, engine="openpyxl") as writer:
                template_df.to_excel(writer, index=False, sheet_name="Projects")
                max_len = max(len(wards), len(subcounties), len(departments), 5)
                ref_df = pd.DataFrame({
                    "Valid Subcounties": sorted(subcounties['name'].tolist())
                        + [""] * (max_len - len(subcounties)),
                    "Valid Wards": sorted(wards['name'].tolist())
                        + [""] * (max_len - len(wards)),
                    "Valid Departments": sorted(departments['name'].tolist())
                        + [""] * (max_len - len(departments)),
                    "Valid Statuses": [
                        "Not Started", "In Progress", "Completed",
                        "Stalled", "Cancelled",
                    ] + [""] * (max_len - 5),
                })
                ref_df.to_excel(writer, index=False, sheet_name="Reference")
            tcol2.download_button(
                "⬇️ Excel Template",
                data=buf.getvalue(),
                file_name="embu_projects_template.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                use_container_width=True,
            )
        except ImportError:
            tcol2.info("Install `openpyxl` for Excel support.")

        # Word template
        try:
            from docx import Document  # noqa: F401
            doc = Document()
            doc.add_heading("Embu County — Project Import Template", level=1)
            doc.add_paragraph(
                "Fill in one table row per project. Do not rename the headers."
            )
            table = doc.add_table(rows=1, cols=len(EXPECTED_COLUMNS))
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, col in enumerate(EXPECTED_COLUMNS):
                hdr[i].text = col
            # Add a sample row
            sample = [
                "Sample Project Name", "Manyatta", "Kirimari", "Education",
                "15-Jul-24", "30-Jun-25", "1500000", "1499891",
                "100", "Completed", "Sample remarks",
            ]
            row = table.add_row().cells
            for i, val in enumerate(sample):
                row[i].text = val

            buf = io.BytesIO()
            doc.save(buf)
            tcol3.download_button(
                "⬇️ Word Template",
                data=buf.getvalue(),
                file_name="embu_projects_template.docx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                use_container_width=True,
            )
        except ImportError:
            tcol3.info("Install `python-docx` for Word support.")

    st.markdown("---")

    # =====================================================
    # INPUT METHOD SELECTOR
    # =====================================================
    input_method = st.radio(
        "Choose how to provide your data",
        options=[
            "📁 Upload File (Excel / CSV / Word)",
            "📋 Paste Text (TSV / CSV)",
        ],
        horizontal=True,
        key="import_method",
    )

    df = None

    # ---- Method 1: File upload ----
    if input_method.startswith("📁"):
        uploaded = st.file_uploader(
            "Upload a file",
            type=["csv", "xlsx", "xls", "docx"],
            key="import_projects_file",
            help="Accepted formats: .xlsx, .xls, .csv, .docx",
        )

        if uploaded is None:
            st.info("👆 Upload a file to preview and import projects.")
            return

        try:
            df = _read_uploaded_file(uploaded)
        except Exception as e:
            st.error(f"❌ Could not read file: {e}")
            return

    # ---- Method 2: Paste text ----
    else:
        pasted = st.text_area(
            "Paste your data (one row per line, values separated by TAB or comma)",
            height=220,
            placeholder=(
                "Project Name\tSubcounty\tWard\tDepartment\t"
                "Commencement Date\tExpected Completion Date\t"
                "Total Funding\tTotal Expenditure\t% Completion\tStatus\tRemarks\n"
                "Sample Project\tManyatta\tKirimari\tEducation\t"
                "15-Jul-24\t30-Jun-25\t1500000\t1499891\t100\tCompleted\tDone"
            ),
            key="import_paste",
        )
        if not pasted.strip():
            st.info("👆 Paste your data to preview and import.")
            return

        # Detect separator
        first_line = pasted.strip().splitlines()[0]
        sep = "\t" if "\t" in first_line else ","
        try:
            df = pd.read_csv(
                io.StringIO(pasted), sep=sep, engine="python"
            )
        except Exception as e:
            st.error(f"❌ Could not parse pasted text: {e}")
            return

    # =====================================================
    # PREVIEW + VALIDATION
    # =====================================================
    if df is None or df.empty:
        st.warning("No data found in the provided input.")
        return

    df.columns = [str(c).strip() for c in df.columns]

    required = {"Project Name", "Subcounty", "Ward", "Department"}
    missing = required - set(df.columns)
    if missing:
        st.error(
            f"❌ Missing required columns: {', '.join(sorted(missing))}. "
            "Download a template for the correct format."
        )
        return

    st.markdown(f"**Preview — {len(df)} rows loaded**")
    st.dataframe(df.head(10), use_container_width=True, hide_index=True)

    # ---- Reference lookups ----
    sc_lookup = {
        n.lower(): i for i, n in zip(subcounties['id'], subcounties['name'])
    }
    ward_lookup = {
        (n.lower(), sc): i
        for i, n, sc in zip(wards['id'], wards['name'], wards['subcounty_id'])
    }
    dept_lookup = {
        n.lower(): i for i, n in zip(departments['id'], departments['name'])
    }
    valid_statuses = {
        "not started", "in progress", "completed", "stalled", "cancelled"
    }

    errors = []
    parsed_rows = []

    for idx, row in df.iterrows():
        row_num = idx + 2  # +2 = header + zero-index
        row_errors = []

        pname = _clean_str(row.get("Project Name"))
        if not pname:
            row_errors.append("Project Name is required")

        sc_name = _clean_str(row.get("Subcounty"))
        sc_id = None
        if not sc_name:
            row_errors.append("Subcounty is required")
        else:
            sc_id = sc_lookup.get(sc_name.lower())
            if sc_id is None:
                row_errors.append(f"Unknown Subcounty '{sc_name}'")

        ward_name = _clean_str(row.get("Ward"))
        ward_id = None
        if not ward_name:
            row_errors.append("Ward is required")
        elif sc_id is not None:
            ward_id = ward_lookup.get((ward_name.lower(), sc_id))
            if ward_id is None:
                row_errors.append(f"Ward '{ward_name}' not in '{sc_name}'")

        dept_name = _clean_str(row.get("Department"))
        dept_id = None
        if not dept_name:
            row_errors.append("Department is required")
        else:
            dept_id = dept_lookup.get(dept_name.lower())
            if dept_id is None:
                row_errors.append(f"Unknown Department '{dept_name}'")

        # Dates
        comm_raw = row.get("Commencement Date")
        exp_raw = row.get("Expected Completion Date")
        comm = _parse_date(comm_raw)
        exp = _parse_date(exp_raw)
        if _clean_str(comm_raw) and comm is None:
            row_errors.append(f"Invalid Commencement Date '{comm_raw}'")
        if _clean_str(exp_raw) and exp is None:
            row_errors.append(f"Invalid Expected Completion Date '{exp_raw}'")
        if comm and exp and exp < comm:
            row_errors.append("Expected Completion before Commencement")

        # Numbers
        funding = _parse_number(row.get("Total Funding"))
        expenditure = _parse_number(row.get("Total Expenditure"))
        if funding is None:
            row_errors.append("Invalid Total Funding")
            funding = 0
        if expenditure is None:
            row_errors.append("Invalid Total Expenditure")
            expenditure = 0
        if funding > 0 and expenditure > funding:
            row_errors.append("Expenditure exceeds Funding")

        pct = _parse_number(row.get("Percentage Completion")) or 0
        pct = max(0.0, min(100.0, float(pct)))

        # Status
        status_raw = _clean_str(row.get("Status"))
        if not status_raw:
            status = "Not Started"
        elif status_raw.lower() not in valid_statuses:
            row_errors.append(f"Invalid Status '{status_raw}'")
            status = "Not Started"
        else:
            status = status_raw.title()

        remarks = _clean_str(row.get("Remarks"))

        if row_errors:
            errors.append({"Row": row_num, "Errors": "; ".join(row_errors)})

        parsed_rows.append({
            "row_num": row_num,
            "project_name": pname,
            "subcounty_id": sc_id,
            "ward_id": ward_id,
            "department_id": dept_id,
            "commencement_date": comm,
            "expected_completion_date": exp,
            "total_funding": funding,
            "total_expenditure": expenditure,
            "percentage_completion": pct,
            "project_status": status,
            "remarks": remarks,
            "has_error": bool(row_errors),
        })

    # ---- Validation summary ----
    st.markdown("---")
    st.markdown("### 🔍 Validation")

    error_count = len(errors)
    ok_count = len(parsed_rows) - error_count

    m1, m2, m3 = st.columns(3)
    m1.metric("Total Rows", len(parsed_rows))
    m2.metric("✅ Ready", ok_count)
    m3.metric("❌ With Errors", error_count)

    if errors:
        with st.expander(f"⚠️ {error_count} row(s) have errors", expanded=True):
            st.dataframe(
                pd.DataFrame(errors), use_container_width=True, hide_index=True
            )

    if ok_count == 0:
        st.error("No valid rows to import.")
        return

    # ---- Import ----
    st.markdown("---")
    st.markdown("### 🚀 Import")

    confirm = st.checkbox(
        f"I confirm I want to import {ok_count} project(s) into the database.",
        key="import_confirm",
    )

    if st.button(
        "📥 Import Projects",
        type="primary",
        disabled=not confirm,
        use_container_width=True,
        key="do_import_btn",
    ):
        inserted = 0
        failed = 0
        failures = []
        rows_to_import = [r for r in parsed_rows if not r["has_error"]]

        progress = st.progress(0, text="Importing...")

        for i, r in enumerate(rows_to_import):
            ok, msg = execute_write(
                """
                INSERT INTO projects
                    (project_name, subcounty_id, ward_id, department_id,
                     commencement_date, expected_completion_date,
                     total_funding, total_expenditure,
                     percentage_completion, project_status, remarks)
                VALUES
                    (:name, :sc, :ward, :dept, :comm, :exp,
                     :fund, :spent, :pct, :status, :remarks)
                """,
                {
                    "name": r["project_name"],
                    "sc": r["subcounty_id"],
                    "ward": r["ward_id"],
                    "dept": r["department_id"],
                    "comm": r["commencement_date"],
                    "exp": r["expected_completion_date"],
                    "fund": r["total_funding"],
                    "spent": r["total_expenditure"],
                    "pct": r["percentage_completion"],
                    "status": r["project_status"],
                    "remarks": r["remarks"],
                },
            )
            if ok:
                inserted += 1
            else:
                failed += 1
                failures.append({"Row": r["row_num"], "Error": msg})

            progress.progress(
                (i + 1) / len(rows_to_import),
                text=f"Importing... {i + 1}/{len(rows_to_import)}",
            )

        progress.empty()

        if inserted:
            st.success(f"✅ Successfully imported {inserted} project(s).")
            st.cache_data.clear()
        if failed:
            st.error(f"❌ {failed} row(s) failed.")
            if failures:
                st.dataframe(
                    pd.DataFrame(failures), use_container_width=True,
                    hide_index=True,
                )
        if inserted and not failed:
            st.balloons()
            st.info("Switch to **View Projects** to see your imported data.")


# =====================================================
# FILE READERS
# =====================================================
def _read_uploaded_file(uploaded):
    """Dispatch to the right reader based on file extension."""
    name = uploaded.name.lower()

    if name.endswith(".csv"):
        return pd.read_csv(uploaded)

    if name.endswith((".xlsx", ".xls")):
        return pd.read_excel(uploaded)

    if name.endswith(".docx"):
        return _read_word_document(uploaded)

    raise ValueError(f"Unsupported file type: {uploaded.name}")


def _read_word_document(uploaded):
    """
    Extract a DataFrame from a Word document.

    Strategy:
        1. If the doc contains tables, read the first non-empty table
           (row 1 = headers).
        2. Otherwise, parse paragraphs formatted as
           "Column: Value" and assemble one project from them.
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "python-docx is required for Word import. "
            "Add `python-docx>=1.1.0` to requirements.txt."
        )

    doc = Document(uploaded)

    # ---- Strategy 1: Tables ----
    for table in doc.tables:
        if len(table.rows) < 2:
            continue

        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if not any(headers):
            continue

        data = []
        for row in table.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                data.append(values)

        if data:
            # Pad rows to match header length
            n = len(headers)
            data = [(r + [""] * n)[:n] for r in data]
            return pd.DataFrame(data, columns=headers)

    # ---- Strategy 2: Paragraphs with "Column: Value" ----
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if not text:
        raise ValueError(
            "The Word document appears to be empty. "
            "Add a table with the expected columns."
        )

    # Try to detect blocks separated by blank lines (each = one project)
    blocks = re.split(r"\n\s*\n", text.strip())
    rows = []
    for block in blocks:
        record = {}
        for line in block.splitlines():
            if ":" in line:
                key, _, val = line.partition(":")
                record[key.strip()] = val.strip()
        if record:
            rows.append(record)

    if not rows:
        raise ValueError(
            "Could not extract project data from the Word document. "
            "Use the Word template with a table, or format lines as "
            "'Column: Value'."
        )

    return pd.DataFrame(rows)


# =====================================================
# UPDATE PROJECT
# =====================================================
def _render_edit_form(projects_df, subcounties, wards, departments):
    edit_id = st.selectbox(
        "Select project to edit",
        options=projects_df['id'].tolist(),
        format_func=lambda x: f"#{x} — "
        f"{projects_df.loc[projects_df['id'] == x, 'project_name'].values[0]}",
        key="edit_project_select",
    )

    current = projects_df[projects_df['id'] == edit_id].iloc[0]

    with st.form("edit_project"):
        col1, col2 = st.columns(2)
        new_name = col1.text_input("Project Name", value=current['project_name'])

        sc_ids = subcounties['id'].tolist()
        sc_names = subcounties['name'].tolist()
        current_sc_name = current['subcounty']
        sc_index = sc_names.index(current_sc_name) if current_sc_name in sc_names else 0
        new_sc_id = col2.selectbox(
            "Subcounty", options=sc_ids, index=sc_index,
            format_func=lambda x: subcounties.loc[
                subcounties['id'] == x, 'name'
            ].values[0],
            key="edit_sc",
        )

        ward_subset = wards[wards['subcounty_id'] == new_sc_id]
        ward_ids = ward_subset['id'].tolist()
        ward_names = ward_subset['name'].tolist()
        current_ward_name = current['ward']
        ward_index = ward_names.index(current_ward_name) \
            if current_ward_name in ward_names else 0
        new_ward_id = col1.selectbox(
            "Ward", options=ward_ids, index=ward_index,
            format_func=lambda x: ward_subset.loc[
                ward_subset['id'] == x, 'name'
            ].values[0],
            key="edit_ward",
        ) if ward_ids else None

        dept_ids = departments['id'].tolist()
        dept_names = departments['name'].tolist()
        current_dept_name = current['department']
        dept_index = dept_names.index(current_dept_name) \
            if current_dept_name in dept_names else 0
        new_dept_id = col2.selectbox(
            "Department", options=dept_ids, index=dept_index,
            format_func=lambda x: departments.loc[
                departments['id'] == x, 'name'
            ].values[0],
            key="edit_dept",
        )

        col3, col4 = st.columns(2)
        new_commencement = col3.date_input(
            "Project Commencement Date",
            value=_to_date(current['commencement_date']),
        )
        new_expected = col4.date_input(
            "Expected Date of Completion",
            value=_to_date(current['expected_completion_date']),
        )

        col5, col6 = st.columns(2)
        new_funding = col5.number_input(
            "Total Funding", min_value=0.0, step=1000.0,
            value=float(current['total_funding'] or 0), format="%.2f",
        )
        new_expenditure = col6.number_input(
            "Total Expenditure", min_value=0.0, step=1000.0,
            value=float(current['total_expenditure'] or 0), format="%.2f",
        )

        col7, col8 = st.columns(2)
        new_pct = col7.slider(
            "Percentage (%) of Completion",
            min_value=0, max_value=100,
            value=int(current['percentage_completion'] or 0), step=1,
        )
        statuses = ["Not Started", "In Progress", "Completed", "Stalled", "Cancelled"]
        status_index = statuses.index(current['project_status']) \
            if current['project_status'] in statuses else 0
        new_status = col8.selectbox(
            "Status", statuses, index=status_index, key="edit_status"
        )

        new_remarks = st.text_area(
            "Remarks", value=current['remarks'] or "", key="edit_remarks"
        )

        if st.form_submit_button("💾 Update Project", use_container_width=True):
            if not new_name.strip() or not new_ward_id:
                st.error("Project Name and Ward are required.")
            elif new_expenditure > new_funding and new_funding > 0:
                st.error("Total Expenditure cannot exceed Total Funding.")
            elif new_commencement and new_expected and new_expected < new_commencement:
                st.error("Expected Completion cannot be before Commencement.")
            else:
                ok, msg = execute_write(
                    """
                    UPDATE projects
                    SET project_name = :name,
                        subcounty_id = :sc,
                        ward_id = :ward,
                        department_id = :dept,
                        project_status = :status,
                        remarks = :remarks,
                        commencement_date = :comm,
                        expected_completion_date = :exp,
                        total_funding = :fund,
                        total_expenditure = :spent,
                        percentage_completion = :pct,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = :id
                    """,
                    {
                        "name": new_name.strip(),
                        "sc": new_sc_id,
                        "ward": new_ward_id,
                        "dept": new_dept_id,
                        "status": new_status,
                        "remarks": new_remarks.strip() or None,
                        "comm": new_commencement,
                        "exp": new_expected,
                        "fund": new_funding,
                        "spent": new_expenditure,
                        "pct": float(new_pct),
                        "id": edit_id,
                    },
                )
                if ok:
                    st.success(f"✅ Project #{edit_id} updated.")
                    st.cache_data.clear()
                    st.rerun()
                else:
                    st.error(f"❌ {msg}")


# =====================================================
# DELETE PROJECT
# =====================================================
def _render_delete_form(projects_df):
    del_id = st.selectbox(
        "Select project to delete",
        options=projects_df['id'].tolist(),
        format_func=lambda x: f"#{x} — "
        f"{projects_df.loc[projects_df['id'] == x, 'project_name'].values[0]}",
        key="delete_project_select",
    )

    confirm = st.checkbox(
        f"I confirm I want to permanently delete project #{del_id}",
        key="delete_confirm",
    )

    if st.button(
        "Delete Project", type="primary",
        disabled=not confirm, use_container_width=True,
    ):
        ok, msg = execute_write(
            "DELETE FROM projects WHERE id = :id", {"id": del_id}
        )
        if ok:
            st.success(f"✅ Project #{del_id} deleted.")
            st.cache_data.clear()
            st.rerun()
        else:
            st.error(f"❌ {msg}")
